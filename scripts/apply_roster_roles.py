"""Aplica a proposta de elenco no Excel e gera um relatorio auditavel.

Uso: python scripts/apply_roster_roles.py outputs/aplicacao-elencos-2026-10-21.json
"""

import json
import os
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
PLAN_PATH = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "outputs/aplicacao-elencos-2026-10-21.json")
PLAN = json.loads(PLAN_PATH.read_text(encoding="utf-8"))
TEAM_PATH = ROOT / "dados_excel/teams.xlsx"
PLAYER_PATH = ROOT / "dados_excel/players.xlsx"
OUT_DIR = ROOT / "outputs" / f"aplicacao-elencos-{PLAN['asOf']}"
REPORT_PATH = OUT_DIR / "relatorio.xlsx"


def headers(sheet):
    return {cell.value: cell.column for cell in sheet[1] if cell.value}


def trim(value):
    return str(value).strip() if value is not None else ""


teams_book = load_workbook(TEAM_PATH)
teams_sheet = teams_book.active
team_headers = headers(teams_sheet)
players_book = load_workbook(PLAYER_PATH)
players_sheet = players_book.active
player_headers = headers(players_sheet)
team_rows = {trim(teams_sheet.cell(row, team_headers["team"]).value).lower(): row for row in range(2, teams_sheet.max_row + 1)}
player_rows = {}
for row in range(2, players_sheet.max_row + 1):
    player_id = trim(players_sheet.cell(row, player_headers["puuid"]).value) or trim(players_sheet.cell(row, player_headers["Jogador"]).value)
    if player_id in player_rows:
        raise ValueError(f"ID duplicado em players.xlsx: {player_id}")
    player_rows[player_id] = row

# Todas as verificacoes ocorrem antes de modificar os arquivos.
if len(team_rows) != teams_sheet.max_row - 1:
    raise ValueError("Equipes duplicadas em teams.xlsx")
if not set(team_rows).issubset({team["id"] for team in PLAN["teams"]}):
    raise ValueError("Equipes da proposta nao correspondem a teams.xlsx")
for team in PLAN["teams"]:
    if team["id"] not in team_rows:
        if team["decisions"]:
            raise ValueError(f"Equipe com elenco sem linha em teams.xlsx: {team['id']}")
        continue
    row = team_rows[team["id"]]
    if not team["eligible"] and any(item["decision"] != "Ignorar" for item in team["decisions"]):
        raise ValueError(f"Equipe ignorada com decisao ativa: {team['id']}")
    for item in team["decisions"]:
        slot = item["slot"]
        if trim(teams_sheet.cell(row, team_headers[f"player{slot}"]).value) != item["player"]:
            raise ValueError(f"Nome mudou em teams.xlsx: {team['id']} player{slot}")
        if trim(teams_sheet.cell(row, team_headers[f"role{slot}"]).value):
            raise ValueError(f"Funcao ja preenchida em teams.xlsx: {team['id']} role{slot}")
        player_row = player_rows.get(item["playerId"])
        if not player_row or trim(players_sheet.cell(player_row, player_headers["Jogador"]).value) != item["player"]:
            raise ValueError(f"Perfil do jogador nao corresponde: {item['player']}")

OUT_DIR.mkdir(parents=True, exist_ok=True)
shutil.copy2(TEAM_PATH, OUT_DIR / "teams-original.xlsx")
shutil.copy2(PLAYER_PATH, OUT_DIR / "players-original.xlsx")

cleared_current = 0
for team in PLAN["teams"]:
    if not team["eligible"] or team["id"] not in team_rows:
        continue
    row = team_rows[team["id"]]
    retained = [item for item in team["decisions"] if item["decision"] != "Remover"]
    for slot in range(1, 15):
        teams_sheet.cell(row, team_headers[f"player{slot}"]).value = None
        teams_sheet.cell(row, team_headers[f"role{slot}"]).value = None
    for slot, item in enumerate(retained, 1):
        teams_sheet.cell(row, team_headers[f"player{slot}"], item["player"])
        teams_sheet.cell(row, team_headers[f"role{slot}"], item["decision"])
    for item in team["decisions"]:
        if item["decision"] != "Remover":
            continue
        player_row = player_rows[item["playerId"]]
        current_cell = players_sheet.cell(player_row, player_headers["current_team"])
        if trim(current_cell.value) == team["id"]:
            current_cell.value = None
            cleared_current += 1


def save_atomically(book, destination):
    temp = destination.with_name(f".{destination.stem}-novo.xlsx")
    book.save(temp)
    os.replace(temp, destination)


save_atomically(teams_book, TEAM_PATH)
save_atomically(players_book, PLAYER_PATH)

report = Workbook()
summary_sheet = report.active
summary_sheet.title = "Resumo por equipe"
summary_sheet.append(["Equipe", "ID", "Series em 60 dias", "Decisao", "Antes", "Titulares", "Reservas", "Removidos", "Depois"])
details_sheet = report.create_sheet("Decisoes por jogador")
details_sheet.append([
    "Equipe", "ID equipe", "Jogador", "Slot anterior", "Decisao", "Motivo", "Top 5", "Series ultimas 5",
    "Series pela equipe", "Mapas pela equipe", "Ultima atuacao", "Dias sem jogar", "Series perdidas",
    "Datas perdidas", "Series por outra equipe apos", "Outra equipe", "current_team anterior",
])
for team in PLAN["teams"]:
    counts = Counter(item["decision"] for item in team["decisions"])
    summary_sheet.append([
        team["name"], team["id"], team["recentSeries60"], "Aplicar" if team["eligible"] else "Ignorar",
        len(team["decisions"]), counts["Titular"], counts["Reserva"], counts["Remover"],
        len(team["decisions"]) - counts["Remover"],
    ])
    for item in team["decisions"]:
        details_sheet.append([
            team["name"], team["id"], item["player"], item["slot"], item["decision"], item["reason"],
            "Sim" if item["topFive"] else "Nao", item["lastFivePlayed"], item["ownSeries"], item["ownMaps"],
            item["ownLast"], item["daysSinceOwn"], item["missedSeries"], item["missedDates"],
            item["otherSeriesAfter"], item["otherTeam"], item["currentTeam"],
        ])
for sheet, widths in [
    (summary_sheet, [31, 30, 20, 16, 10, 13, 12, 13, 10]),
    (details_sheet, [31, 30, 24, 16, 14, 65, 10, 17, 18, 17, 17, 16, 16, 14, 20, 25, 25]),
]:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="183153")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
report.save(REPORT_PATH)
print(json.dumps({"report": str(REPORT_PATH), "summary": PLAN["summary"], "currentTeamCleared": cleared_current}, ensure_ascii=False))
