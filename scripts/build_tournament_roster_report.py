"""Gera um controle de investigação por campeonato sem alterar as fontes."""

import argparse
import json
import subprocess
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
LPE_URL = "https://www.challengermode.com/s/Thinkgameshub/tournaments/a880743a-3f64-461f-33ef-08de95af1b37/participants"

NODE_SCRIPT = r"""
const fs = require('fs');
const { decode } = require('./db-codec.js');
const db = decode(JSON.parse(fs.readFileSync('database.json', 'utf8')));
const events = (db.tournaments || []).map(({ id, name, startAt, endAt, teams, matches, mapGames }) =>
  ({ id, name, startAt, endAt, teams: teams || [], matches, mapGames }));
const observed = new Map();
for (const series of db.matchSeries || []) {
  const seen = new Set();
  for (const match of series.maps || []) {
    for (const player of match.players || []) {
      if (!player.teamId || !player.id) continue;
      const key = [series.eventId, player.teamId, player.id].join('|');
      const row = observed.get(key) || { eventId: series.eventId, teamId: player.teamId,
        playerId: player.id, nick: player.nick || '', maps: 0, series: 0, firstAt: 0, lastAt: 0 };
      row.maps++;
      row.nick = player.nick || row.nick;
      const at = Number(match.startedAt || series.startedAt || 0);
      row.firstAt = row.firstAt ? Math.min(row.firstAt, at) : at;
      row.lastAt = Math.max(row.lastAt, at);
      if (!seen.has(key)) { row.series++; seen.add(key); }
      observed.set(key, row);
    }
  }
}
process.stdout.write(JSON.stringify({ events, observed: [...observed.values()] }));
"""


def build(output):
    if output.exists():
        raise SystemExit(f"Arquivo já existe; preserve as anotações: {output}")
    result = subprocess.run(["node", "-e", NODE_SCRIPT], cwd=ROOT, check=True, capture_output=True, text=True, encoding="utf-8")
    data = json.loads(result.stdout)
    events = sorted(data["events"], key=lambda row: (row.get("startAt") or "", row["name"]))
    observed = sorted(data["observed"], key=lambda row: (row["eventId"], row["teamId"], row["nick"].lower()))
    participants = defaultdict(set)
    for row in observed:
        participants[row["eventId"]].add(row["teamId"])

    book = Workbook()
    overview = book.active
    overview.title = "Campeonatos"
    overview.append(["Campeonato", "ID", "Início", "Fim", "Equipes com mapas", "Séries locais", "Mapas locais", "Cadastro externo", "Situação da revisão", "Observação"])
    for event in events:
        is_lpe = event["id"] == "lpe"
        overview.append([
            event["name"], event["id"], (event.get("startAt") or "")[:10], (event.get("endAt") or "")[:10],
            len(participants[event["id"]]), event.get("matches") or 0, event.get("mapGames") or 0,
            LPE_URL if is_lpe else "", "Parcial: CAAP" if is_lpe else "Pendente",
            "Partidas locais não comprovam a lista completa de inscritos.",
        ])
        if is_lpe:
            cell = overview.cell(overview.max_row, 8)
            cell.hyperlink = LPE_URL
            cell.style = "Hyperlink"

    sheet = book.create_sheet("Jogadores observados")
    sheet.append(["Campeonato", "Equipe", "Jogador no jogo", "Mapas", "Séries", "Primeiro mapa", "Último mapa", "Origem"])
    for row in observed:
        sheet.append([row["eventId"], row["teamId"], row["nick"], row["maps"], row["series"],
                      datetime.fromtimestamp(row["firstAt"] / 1000).strftime("%Y-%m-%d") if row["firstAt"] else "",
                      datetime.fromtimestamp(row["lastAt"] / 1000).strftime("%Y-%m-%d") if row["lastAt"] else "",
                      "Partida competitiva local"])

    registration = book.create_sheet("Inscrições verificadas")
    registration.append(["Campeonato", "Equipe", "Jogador na plataforma", "Jogador no site", "Função no torneio", "Fonte", "Verificado em", "Limite da evidência"])
    for platform, local, role in [
        ("Metalvlr", "Metal", "Lineup"), ("lopesg", "Lopes", "Lineup"),
        ("Burrito42", "Burrito", "Lineup"), ("sobrunin", "sobrunin", "Lineup"),
        ("mrzz / coutinho#iwnl", "Coutinho", "Lineup"), ("medinavlr", "Medina", "Substitute"),
    ]:
        registration.append(["LPE", "caap_hellhounds", platform, local, role, LPE_URL,
                             "2026-10-21", "Função na inscrição do LPE; não determina função atual."])
        cell = registration.cell(registration.max_row, 6)
        cell.hyperlink = LPE_URL
        cell.style = "Hyperlink"

    notes = book.create_sheet("Como revisar")
    for row in [
        ["Escopo", "25 campeonatos com arquivos de partidas no projeto; as inscrições externas ainda não foram verificadas integralmente."],
        ["Jogadores observados", "Somente quem aparece em mapa competitivo. Ausência de mapa não significa ausência do cadastro do torneio."],
        ["Lineup", "Função de titular naquele campeonato, conforme a aba Teams do Challengermode."],
        ["Substitute", "Função de substituto naquele campeonato; não define a função atual no time."],
        ["Coach", "Função técnica separada de titular e reserva."],
        ["Próximo passo", "Abrir a página oficial de cada torneio; registrar lineup, substitutos e técnicos de cada equipe; comparar com partidas e com teams.xlsx antes de mudar a função atual."],
    ]:
        notes.append(row)

    widths = {
        "Campeonatos": [37, 29, 14, 14, 20, 15, 15, 60, 22, 65],
        "Jogadores observados": [30, 30, 27, 10, 10, 16, 16, 28],
        "Inscrições verificadas": [16, 28, 30, 24, 18, 60, 18, 65],
        "Como revisar": [25, 105],
    }
    for ws in book:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = ws.dimensions
        for index, width in enumerate(widths[ws.title], start=1):
            ws.column_dimensions[chr(64 + index)].width = width
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="191D29")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 28
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=ws.title == "Como revisar")
        if ws.title != "Como revisar":
            table = Table(displayName={"Campeonatos": "CampeonatosRevisao", "Jogadores observados": "JogadoresObservados", "Inscrições verificadas": "InscricoesVerificadas"}[ws.title], ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    print(json.dumps({"file": str(output), "tournaments": len(events), "observedPlayers": len(observed), "verifiedRegistrations": 6}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=ROOT / "outputs" / f"investigacao-inscricoes-{date.today().isoformat()}.xlsx")
    args = parser.parse_args()
    build(args.out)
